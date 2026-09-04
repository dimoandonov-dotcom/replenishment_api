"""
Административни endpoints - всичко, което досега се правеше със скриптове,
е достъпно и през API-то: магазини, адреси, доставчици, графици, периоди
на затваряне, опаковки, стоп кодове, планограма, min/max.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import select, func, delete
from sqlalchemy.orm import Session

from . import models as m
from .db import get_db
from .imports import normalize_store_name

router = APIRouter(tags=["Администриране"])


class StoreIn(BaseModel):
    name: str
    address: str | None = None
    is_active: bool = True


class StorePatch(BaseModel):
    name: str | None = None
    address: str | None = None
    is_active: bool | None = None


def _store_out(s: m.Store) -> dict:
    return {"id": s.id, "name": s.name, "address": s.address, "is_active": s.is_active}


@router.post("/stores")
def create_store(payload: StoreIn, db: Session = Depends(get_db)):
    if db.execute(select(m.Store).where(m.Store.name == payload.name)).scalar_one_or_none():
        raise HTTPException(409, "Вече има магазин с това име")
    next_id = (db.execute(select(func.max(m.Store.id))).scalar() or 0) + 1
    s = m.Store(id=next_id, name=payload.name, address=payload.address,
                is_active=payload.is_active)
    db.add(s)
    db.flush()
    norm = normalize_store_name(payload.name)
    if norm and not db.execute(
        select(m.StoreAlias).where(m.StoreAlias.alias_normalized == norm)
    ).scalar_one_or_none():
        db.add(m.StoreAlias(store_id=s.id, alias_normalized=norm))
    db.commit()
    return _store_out(s)


@router.patch("/stores/{store_id}")
def update_store(store_id: int, payload: StorePatch, db: Session = Depends(get_db)):
    s = db.get(m.Store, store_id)
    if not s:
        raise HTTPException(404, "Няма такъв магазин")
    if payload.name is not None:
        s.name = payload.name
    if payload.address is not None:
        s.address = payload.address
    if payload.is_active is not None:
        s.is_active = payload.is_active
    db.commit()
    return _store_out(s)


@router.post("/stores/{store_id}/close")
def close_store(store_id: int, db: Session = Depends(get_db)):
    s = db.get(m.Store, store_id)
    if not s:
        raise HTTPException(404, "Няма такъв магазин")
    s.is_active = False
    db.commit()
    return {"id": s.id, "name": s.name, "is_active": False,
            "note": "Обектът вече не участва в автоматични заявки"}


@router.post("/stores/{store_id}/reopen")
def reopen_store(store_id: int, db: Session = Depends(get_db)):
    s = db.get(m.Store, store_id)
    if not s:
        raise HTTPException(404, "Няма такъв магазин")
    s.is_active = True
    db.commit()
    return {"id": s.id, "name": s.name, "is_active": True}


class AliasIn(BaseModel):
    alias: str


@router.post("/stores/{store_id}/aliases")
def add_alias(store_id: int, payload: AliasIn, db: Session = Depends(get_db)):
    if not db.get(m.Store, store_id):
        raise HTTPException(404, "Няма такъв магазин")
    norm = normalize_store_name(payload.alias)
    if not norm:
        raise HTTPException(400, "Празен псевдоним")
    existing = db.execute(
        select(m.StoreAlias).where(m.StoreAlias.alias_normalized == norm)
    ).scalar_one_or_none()
    if existing:
        if existing.store_id != store_id:
            raise HTTPException(409, f"Псевдонимът вече сочи към магазин {existing.store_id}")
        return {"alias_normalized": norm, "store_id": store_id, "note": "вече съществува"}
    db.add(m.StoreAlias(store_id=store_id, alias_normalized=norm))
    db.commit()
    return {"alias_normalized": norm, "store_id": store_id}


@router.get("/stores/{store_id}/aliases")
def list_aliases(store_id: int, db: Session = Depends(get_db)):
    rows = db.execute(
        select(m.StoreAlias).where(m.StoreAlias.store_id == store_id)
    ).scalars().all()
    return [a.alias_normalized for a in rows]


class SupplierIn(BaseModel):
    name: str
    contact_email: str | None = None
    order_format: str = "email"
    replenishment_mode: str = "below_min"
    apply_weekend_buffer: bool = False
    is_active: bool = True


class SupplierPatch(BaseModel):
    name: str | None = None
    contact_email: str | None = None
    order_format: str | None = None
    replenishment_mode: str | None = None
    apply_weekend_buffer: bool | None = None
    is_active: bool | None = None


@router.post("/suppliers")
def create_supplier(payload: SupplierIn, db: Session = Depends(get_db)):
    if db.execute(select(m.Supplier).where(m.Supplier.name == payload.name)).scalar_one_or_none():
        raise HTTPException(409, "Вече има доставчик с това име")
    s = m.Supplier(**payload.model_dump())
    db.add(s)
    db.commit()
    return {"id": s.id, "name": s.name, "contact_email": s.contact_email}


@router.patch("/suppliers/{supplier_id}")
def update_supplier(supplier_id: int, payload: SupplierPatch, db: Session = Depends(get_db)):
    s = db.get(m.Supplier, supplier_id)
    if not s:
        raise HTTPException(404, "Няма такъв доставчик")
    data = payload.model_dump(exclude_none=True)
    if data.get("replenishment_mode") not in (None, "below_min", "daily_topup"):
        raise HTTPException(400, "replenishment_mode: below_min или daily_topup")
    for k, v in data.items():
        setattr(s, k, v)
    db.commit()
    return {"id": s.id, "name": s.name, "contact_email": s.contact_email,
            "replenishment_mode": s.replenishment_mode,
            "apply_weekend_buffer": s.apply_weekend_buffer, "is_active": s.is_active}


class ArticlePatch(BaseModel):
    name: str | None = None
    pack_size: int | None = Field(None, ge=1)
    pack_type: str | None = None
    is_active: bool | None = None
    default_supplier_id: int | None = None


@router.patch("/articles/{sku}")
def update_article(sku: str, payload: ArticlePatch, db: Session = Depends(get_db)):
    a = db.execute(select(m.Article).where(m.Article.sku == sku)).scalar_one_or_none()
    if not a:
        raise HTTPException(404, f"Няма артикул с SKU {sku}")
    for k, v in payload.model_dump(exclude_none=True).items():
        setattr(a, k, v)
    db.commit()
    return {"sku": a.sku, "name": a.name, "pack_size": a.pack_size,
            "pack_type": a.pack_type, "is_active": a.is_active}


class StopCodesIn(BaseModel):
    skus: list[str]
    blocked: bool = True


@router.post("/articles/stop-codes")
def set_stop_codes(payload: StopCodesIn, db: Session = Depends(get_db)):
    arts = db.execute(
        select(m.Article).where(m.Article.sku.in_(payload.skus))
    ).scalars().all()
    found = {a.sku for a in arts}
    removed_settings = 0
    for a in arts:
        a.is_active = not payload.blocked
        if payload.blocked:
            removed_settings += db.execute(
                delete(m.StoreArticleSetting).where(
                    m.StoreArticleSetting.article_id == a.id)
            ).rowcount or 0
    db.commit()
    return {"updated": len(arts), "blocked": payload.blocked,
            "settings_removed": removed_settings,
            "unknown_skus": sorted(set(payload.skus) - found)}


@router.get("/articles/stop-codes")
def list_stop_codes(db: Session = Depends(get_db)):
    rows = db.execute(
        select(m.Article).where(m.Article.is_active.is_(False)).order_by(m.Article.sku)
    ).scalars().all()
    return [{"sku": a.sku, "name": a.name} for a in rows]


class ScheduleIn(BaseModel):
    store_id: int
    supplier_id: int
    order_weekday: int = Field(ge=1, le=7)
    delivery_weekday: int | None = Field(None, ge=1, le=7)
    order_cutoff_time: str | None = None
    confirmed: bool = False


@router.get("/schedules")
def list_schedules(store_id: int | None = None, supplier_id: int | None = None,
                   db: Session = Depends(get_db)):
    stmt = select(m.StoreSupplierSchedule)
    if store_id:
        stmt = stmt.where(m.StoreSupplierSchedule.store_id == store_id)
    if supplier_id:
        stmt = stmt.where(m.StoreSupplierSchedule.supplier_id == supplier_id)
    rows = db.execute(stmt).scalars().all()
    return [{"id": r.id, "store_id": r.store_id, "supplier_id": r.supplier_id,
             "order_weekday": r.weekday, "delivery_weekday": r.delivery_weekday,
             "order_cutoff_time": str(r.order_cutoff_time) if r.order_cutoff_time else None,
             "confirmed": r.confirmed} for r in rows]


@router.put("/schedules")
def upsert_schedules(items: list[ScheduleIn], db: Session = Depends(get_db)):
    from datetime import time as _time
    saved = 0
    for it in items:
        cutoff = None
        if it.order_cutoff_time:
            hh, _, mm = it.order_cutoff_time.partition(":")
            cutoff = _time(int(hh), int(mm or 0))
        row = db.execute(
            select(m.StoreSupplierSchedule).where(
                m.StoreSupplierSchedule.store_id == it.store_id,
                m.StoreSupplierSchedule.supplier_id == it.supplier_id,
                m.StoreSupplierSchedule.weekday == it.order_weekday,
            )
        ).scalar_one_or_none()
        if row:
            row.delivery_weekday = it.delivery_weekday
            row.order_cutoff_time = cutoff
            row.confirmed = it.confirmed
        else:
            db.add(m.StoreSupplierSchedule(
                store_id=it.store_id, supplier_id=it.supplier_id,
                weekday=it.order_weekday, delivery_weekday=it.delivery_weekday,
                order_cutoff_time=cutoff, confirmed=it.confirmed))
        saved += 1
    db.commit()
    return {"saved": saved}


@router.delete("/schedules/{schedule_id}")
def delete_schedule(schedule_id: int, db: Session = Depends(get_db)):
    row = db.get(m.StoreSupplierSchedule, schedule_id)
    if not row:
        raise HTTPException(404, "Няма такъв график")
    db.delete(row)
    db.commit()
    return {"deleted": schedule_id}


class ClosureIn(BaseModel):
    supplier_id: int
    start_date: date
    end_date: date
    reason: str | None = None
    buffer_days: int = 0


@router.get("/closures")
def list_closures(supplier_id: int | None = None, db: Session = Depends(get_db)):
    stmt = select(m.SupplierClosure).order_by(m.SupplierClosure.start_date)
    if supplier_id:
        stmt = stmt.where(m.SupplierClosure.supplier_id == supplier_id)
    rows = db.execute(stmt).scalars().all()
    return [{"id": r.id, "supplier_id": r.supplier_id, "start_date": r.start_date,
             "end_date": r.end_date, "reason": r.reason,
             "buffer_days": r.buffer_days} for r in rows]


@router.post("/closures")
def create_closure(payload: ClosureIn, db: Session = Depends(get_db)):
    if payload.end_date < payload.start_date:
        raise HTTPException(400, "Крайната дата е преди началната")
    if not db.get(m.Supplier, payload.supplier_id):
        raise HTTPException(404, "Няма такъв доставчик")
    c = m.SupplierClosure(**payload.model_dump())
    db.add(c)
    db.commit()
    return {"id": c.id, "supplier_id": c.supplier_id,
            "start_date": c.start_date, "end_date": c.end_date, "reason": c.reason}


@router.delete("/closures/{closure_id}")
def delete_closure(closure_id: int, db: Session = Depends(get_db)):
    c = db.get(m.SupplierClosure, closure_id)
    if not c:
        raise HTTPException(404, "Няма такъв период")
    db.delete(c)
    db.commit()
    return {"deleted": closure_id}


def _load_sheet(content: bytes, sheet: str | None):
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    if sheet:
        if sheet not in wb.sheetnames:
            raise HTTPException(400, f"Няма лист '{sheet}'. Налични: {wb.sheetnames}")
        return wb[sheet]
    return wb.worksheets[0]


@router.post("/import/nomenclature")
async def import_nomenclature(
    file: UploadFile = File(...),
    supplier_id: int = Query(..., description="За кой доставчик е номенклатурата"),
    sheet: str | None = Query(None),
    db: Session = Depends(get_db),
):
    if not db.get(m.Supplier, supplier_id):
        raise HTTPException(404, "Няма такъв доставчик")
    ws = _load_sheet(await file.read(), sheet)
    existing = {a.sku: a for a in db.execute(select(m.Article)).scalars().all()}
    created = updated = 0
    for row in ws.iter_rows(values_only=True):
        if not row or not isinstance(row[0], (int, float)):
            continue
        sku, name = str(int(row[0])), str(row[1] or "").strip()
        if not name:
            continue
        a = existing.get(sku)
        if a:
            a.name = name
            a.default_supplier_id = supplier_id
            updated += 1
        else:
            a = m.Article(sku=sku, name=name, default_supplier_id=supplier_id, pack_size=1)
            db.add(a)
            db.flush()
            existing[sku] = a
            created += 1
    db.commit()
    return {"created": created, "updated": updated}


@router.post("/import/pack-sizes")
async def import_pack_sizes(
    file: UploadFile = File(...),
    sheet: str | None = Query(None),
    db: Session = Depends(get_db),
):
    ws = _load_sheet(await file.read(), sheet)
    arts = {a.sku: a for a in db.execute(select(m.Article)).scalars().all()}
    updated = to_units = unknown = 0
    for row in ws.iter_rows(values_only=True):
        if not row or not isinstance(row[0], (int, float)):
            continue
        a = arts.get(str(int(row[0])))
        if not a:
            unknown += 1
            continue
        filled = [(n, int(v)) for n, v in
                  (("Стек", row[2] if len(row) > 2 else None),
                   ("Каса", row[3] if len(row) > 3 else None),
                   ("Кашон", row[4] if len(row) > 4 else None))
                  if isinstance(v, (int, float))]
        if filled:
            a.pack_type, a.pack_size = filled[0]
        else:
            a.pack_type, a.pack_size = "брой", 1
            to_units += 1
        updated += 1
    db.commit()
    return {"updated": updated, "set_to_units": to_units, "unknown_skus": unknown}


@router.post("/import/stop-codes")
async def import_stop_codes(
    file: UploadFile = File(...),
    sheet: str | None = Query(None),
    db: Session = Depends(get_db),
):
    ws = _load_sheet(await file.read(), sheet)
    skus = [str(int(r[0])) for r in ws.iter_rows(values_only=True)
            if r and isinstance(r[0], (int, float))]
    arts = db.execute(select(m.Article).where(m.Article.sku.in_(skus))).scalars().all()
    removed = 0
    for a in arts:
        a.is_active = False
        removed += db.execute(
            delete(m.StoreArticleSetting).where(m.StoreArticleSetting.article_id == a.id)
        ).rowcount or 0
    db.commit()
    return {"blocked": len(arts), "settings_removed": removed,
            "unknown_skus": sorted(set(skus) - {a.sku for a in arts})}


@router.post("/import/planogram")
async def import_planogram(
    file: UploadFile = File(...),
    supplier_id: int = Query(...),
    sheet: str | None = Query(None),
    header_row: int = Query(0, description="Ред със заглавията (0 = първи)"),
    replace: bool = Query(False, description="Изтрий старите настройки на доставчика"),
    db: Session = Depends(get_db),
):
    if not db.get(m.Supplier, supplier_id):
        raise HTTPException(404, "Няма такъв доставчик")
    ws = _load_sheet(await file.read(), sheet)
    rows = list(ws.iter_rows(values_only=True))
    if header_row >= len(rows):
        raise HTTPException(400, "header_row е извън файла")

    lookup: dict[str, int] = {}
    for s in db.execute(select(m.Store)).scalars().all():
        lookup[normalize_store_name(s.name)] = s.id
    for al in db.execute(select(m.StoreAlias)).scalars().all():
        lookup[al.alias_normalized] = al.store_id

    col_store: dict[int, int] = {}
    unmatched: list[str] = []
    for idx, raw in enumerate(rows[header_row]):
        if idx < 2 or not raw:
            continue
        sid = lookup.get(normalize_store_name(str(raw)))
        if sid is None:
            unmatched.append(str(raw))
        else:
            col_store[idx] = sid

    arts = {a.sku: a for a in db.execute(select(m.Article)).scalars().all()}
    placements: set[tuple[int, str]] = set()
    unknown_skus: set[str] = set()
    for r in rows[header_row + 1:]:
        if not r or not isinstance(r[0], (int, float)):
            continue
        sku = str(int(r[0]))
        if sku not in arts:
            unknown_skus.add(sku)
            continue
        for idx, sid in col_store.items():
            if idx < len(r) and r[idx] and str(r[idx]).strip().lower() == "да":
                placements.add((sid, sku))

    if replace:
        old_ids = [a.id for a in arts.values() if a.default_supplier_id == supplier_id]
        if old_ids:
            db.execute(delete(m.StoreArticleSetting).where(
                m.StoreArticleSetting.article_id.in_(old_ids)))
            db.execute(delete(m.SettingsReviewQueue).where(
                m.SettingsReviewQueue.article_id.in_(old_ids)))
            db.flush()

    kept = queued = skipped_blocked = 0
    for store_id, sku in placements:
        a = arts[sku]
        if not a.is_active:
            skipped_blocked += 1
            continue
        if db.get(m.StoreArticleSetting, (store_id, a.id)):
            kept += 1
            continue
        exists_q = db.execute(select(m.SettingsReviewQueue).where(
            m.SettingsReviewQueue.store_id == store_id,
            m.SettingsReviewQueue.article_id == a.id)).scalar_one_or_none()
        if not exists_q:
            db.add(m.SettingsReviewQueue(
                store_id=store_id, article_id=a.id,
                reason="По планограма, но липсват min/max"))
            queued += 1
    db.commit()
    return {"placements": len(placements), "stores_matched": len(col_store),
            "with_existing_minmax": kept, "queued_for_minmax": queued,
            "skipped_blocked": skipped_blocked,
            "unmatched_store_columns": unmatched,
            "unknown_skus": sorted(unknown_skus)[:50]}


@router.post("/import/min-max")
async def import_min_max(
    file: UploadFile = File(...),
    store_id: int | None = Query(None, description="Ако файлът е за един обект"),
    sheet: str | None = Query(None),
    db: Session = Depends(get_db),
):
    def _is_header(row) -> bool:
        if not row:
            return False
        cells = [str(c).strip().upper() for c in row if c]
        return (any(h.startswith("МАТ") for h in cells)
                and any(h.startswith("МИН") for h in cells))

    content = await file.read()
    if sheet:
        ws = _load_sheet(content, sheet)
        rows = list(ws.iter_rows(values_only=True))
    else:
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
        rows = []
        for cand in wb.worksheets:
            r = list(cand.iter_rows(values_only=True))
            if any(_is_header(row) for row in r[:20]):
                rows = r
                break
        if not rows:
            rows = list(wb.worksheets[0].iter_rows(values_only=True))

    hdr_i = None
    for i, r in enumerate(rows[:20]):
        if _is_header(r):
            hdr_i = i
            break

    arts = {a.sku: a for a in db.execute(select(m.Article)).scalars().all()}
    saved = skipped = unknown = invalid = 0

    def _apply(sid: int, sku: str, mn: float, mx: float):
        nonlocal saved, unknown, invalid
        a = arts.get(sku)
        if not a:
            unknown += 1
            return
        if not a.is_active:
            return
        if mx < mn:
            invalid += 1
            return
        row = db.get(m.StoreArticleSetting, (sid, a.id))
        if row:
            row.min_stock, row.max_stock = mn, mx
        else:
            db.add(m.StoreArticleSetting(
                store_id=sid, article_id=a.id, min_stock=mn, max_stock=mx,
                supplier_id=a.default_supplier_id))
        q = db.execute(select(m.SettingsReviewQueue).where(
            m.SettingsReviewQueue.store_id == sid,
            m.SettingsReviewQueue.article_id == a.id)).scalar_one_or_none()
        if q:
            q.resolved = True
        saved += 1

    if hdr_i is not None:
        header = [str(c).strip().upper() if c else "" for c in rows[hdr_i]]
        def col(*names):
            for i, h in enumerate(header):
                if any(n.upper() in h for n in names):
                    return i
            return None
        c_sku, c_min, c_max = col("МАТ"), col("МИН"), col("МАКС", "ОПТ")
        if c_sku is None or c_min is None or c_max is None:
            raise HTTPException(400, f"Не намирам нужните колони. Заглавия: {header}")
        sid = store_id
        if sid is None:
            c_obj = col("ОБЕКТ")
            if c_obj is None:
                raise HTTPException(400, "Подай store_id или добави колона 'Обект'")
        for r in rows[hdr_i + 1:]:
            if not r or not isinstance(r[c_sku], (int, float)):
                continue
            if store_id is None:
                sid = None
                raw = r[c_obj] if c_obj < len(r) else None
                if isinstance(raw, str):
                    lookup = {normalize_store_name(s.name): s.id
                              for s in db.execute(select(m.Store)).scalars().all()}
                    for al in db.execute(select(m.StoreAlias)).scalars().all():
                        lookup[al.alias_normalized] = al.store_id
                    sid = lookup.get(normalize_store_name(raw))
                if sid is None:
                    skipped += 1
                    continue
            mn, mx = r[c_min], r[c_max]
            if mn is None or mx is None:
                skipped += 1
                continue
            _apply(sid, str(int(r[c_sku])), float(mn), float(mx))
    else:
        if store_id is None:
            raise HTTPException(400, "За прост формат е нужен store_id")
        for r in rows:
            if not r or not isinstance(r[0], (int, float)):
                continue
            if len(r) < 3 or r[1] is None or r[2] is None:
                skipped += 1
                continue
            _apply(store_id, str(int(r[0])), float(r[1]), float(r[2]))

    db.commit()
    return {"saved": saved, "skipped_no_values": skipped,
            "unknown_skus": unknown, "invalid_max_below_min": invalid}
